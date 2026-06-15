import sys
import tempfile
import subprocess
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from database import FileDatabase
from indexer import FileIndexer
from searcher import FileSearcher


def write_text_pdf(path: Path, text: str) -> None:
    content = f"BT /F1 12 Tf 72 100 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>\nendobj\n",
        b"4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
        b"5 0 obj\n<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream\nendobj\n",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(data))
        data.extend(obj)
    xref_offset = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    data.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(data)


class ContentSearchTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.managed = self.root / "managed"
        self.managed.mkdir()
        self.db = FileDatabase(self.root / "test.sqlite3")
        self.indexer = FileIndexer(self.db)
        self.searcher = FileSearcher(self.db)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def create_text_file(self, relative: str, content: str) -> Path:
        path = self.managed / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return path

    def scan(self):
        self.indexer.add_directory(self.managed)
        return self.indexer.scan_all()

    def test_content_hits_rank_after_filename_and_path_hits(self) -> None:
        self.create_text_file("预算内容.txt", "文件名命中")
        self.create_text_file("预算内容目录/其他.txt", "路径命中")
        self.create_text_file("普通报告.txt", "这个文件正文包含预算内容，需要被内容搜索命中。")

        stats = self.scan()
        self.assertEqual(stats.content_failed, 0)

        results = self.searcher.search("预算内容")
        self.assertEqual([result.filename for result in results[:3]], ["预算内容.txt", "其他.txt", "普通报告.txt"])
        self.assertEqual([result.match_type for result in results[:3]], ["文件名命中", "路径命中", "内容命中"])
        self.assertEqual([result.rank for result in results[:3]], [1, 4, 5])
        self.assertIn("正文包含预算内容", results[2].snippet)

    def test_extracts_txt_md_csv_docx_and_xlsx_content(self) -> None:
        self.create_text_file("说明.txt", "本地搜索工具")
        self.create_text_file("计划.md", "# 项目计划\n整理建议系统")
        self.create_text_file("明细.csv", "姓名,事项\n张三,付款复核\n")

        docx_path = self.managed / "会议纪要.docx"
        document = Document()
        document.add_paragraph("合同审批需要补充附件")
        document.save(docx_path)

        xlsx_path = self.managed / "奖金表.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "人员"
        sheet.append(["姓名", "事项"])
        sheet.append(["李四", "奖金确认"])
        workbook.save(xlsx_path)

        stats = self.scan()
        self.assertEqual(stats.content_failed, 0)
        self.assertGreaterEqual(stats.content_indexed, 5)

        self.assertEqual(self.searcher.search("本地搜索工具")[0].match_type, "内容命中")
        self.assertEqual(self.searcher.search("整理建议系统")[0].filename, "计划.md")
        self.assertEqual(self.searcher.search("付款复核")[0].filename, "明细.csv")
        self.assertEqual(self.searcher.search("合同审批")[0].filename, "会议纪要.docx")

        excel_result = self.searcher.search("奖金确认")[0]
        self.assertEqual(excel_result.filename, "奖金表.xlsx")
        self.assertIn("[人员 行2]", excel_result.snippet)

    def test_extracts_pdf_content(self) -> None:
        pdf_path = self.managed / "PDF报告.pdf"
        write_text_pdf(pdf_path, "Hello PDF Budget")

        stats = self.scan()
        self.assertEqual(stats.content_failed, 0)

        result = self.searcher.search("budget")[0]
        self.assertEqual(result.filename, "PDF报告.pdf")
        self.assertEqual(result.match_type, "内容命中")
        self.assertIn("[PDF 第1页]", result.snippet)

    def test_pdf_extraction_failure_is_recorded_without_crashing(self) -> None:
        self.create_text_file("坏文件.pdf", "这不是有效 PDF")

        stats = self.scan()
        self.assertEqual(stats.content_failed, 1)

        errors = self.db.fetch_content_errors()
        self.assertEqual(len(errors), 1)
        self.assertTrue(errors[0]["error"])

    def test_office_extraction_timeout_is_recorded_without_hanging(self) -> None:
        self.create_text_file("卡住文件.xlsx", "fake")

        with patch("content_indexer.subprocess.run", side_effect=subprocess.TimeoutExpired("extract", 20)):
            stats = self.scan()

        self.assertEqual(stats.content_failed, 1)
        errors = self.db.fetch_content_errors()
        self.assertEqual(len(errors), 1)
        self.assertIn("内容提取超时", errors[0]["error"])


if __name__ == "__main__":
    unittest.main()
